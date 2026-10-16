"""Fictional crop sales retain their actual totals, review and source history."""

import csv
from decimal import Decimal
from io import BytesIO, StringIO
import json
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from conftest import confirm, fact_data, farm_data


def test_sales_document_csv_review_scope_and_immutable_corrected_output(farmers):
    owner = farmers("Fictional crop seller")
    adviser = farmers("Scoped sales reviewer")
    outsider = farmers("Unrelated financial account")
    farm = owner.create("farms", farm_data("Authored wheat holding"))
    private = owner.create("farms", farm_data("Unshared financial holding"))
    original = (b"FICTIONAL SALES INVOICE - NOT REAL COMMERCE\nDate: 2026-06-11\n"
                b"Buyer: Fictional Mill\nInvoice: AUTHORED-SALE-001\nWheat: 25,50 kg\n"
                b"Printed total: 125,50 EUR\nVAT: unreadable; do not calculate it\n")
    document = owner.post("/documents", expected=201, files={"file": ("authored-sale.txt", original, "text/plain")}, data={"farm_id": farm["id"]}).json()
    extracted = owner.post(f"/documents/{document['id']}/extract", json={"expected_revision": document["revision"], "document_type": "receipt"}).json()
    assert extracted["data"]["extraction"]["transaction_direction"] == "unreviewed"
    assert owner.get("/sales").json()["total"] == 0
    assert owner.get("/purchases").json()["total"] == 0
    assert owner.get("/operations").json()["total"] == 0
    raw = {"farm_id": farm["id"], "title": "Printed wheat sale to Fictional Mill", "category": "crop",
           "date": "2026-06-11", "buyer": "Fictional Mill", "invoice_number": "AUTHORED-SALE-001",
           "quantity": "25,50", "unit": "kg", "amount": "125,50", "currency": "EUR", "credit": False}
    incomplete = owner.post("/workflow/check", json={"kind": "sales", "data": {**raw, "date": None, "unit": None, "amount": None}}).json()
    assert incomplete["can_save"] is False
    assert {issue["field"] for issue in incomplete["issues"] if issue["severity"] == "error"} >= {"date", "unit", "amount"}
    checked = owner.post("/workflow/check", json={"kind": "sales", "data": raw}).json()
    assert checked["can_save"] is True and "decimal_comma" in checked["requires_acknowledgement"]
    assert Decimal(str(checked["normalized_data"]["quantity"])) == Decimal("25.50")
    assert Decimal(str(checked["normalized_data"]["amount"])) == Decimal("125.50")
    assert checked["normalized_data"]["period"] == "2026-06-11"
    # The farmer selects sales explicitly; extraction never chooses direction.
    document_body = {"expected_revision": extracted["revision"], "kind": "sales", "data": raw,
                     "selected_proposal_ids": [], "acknowledged_issue_codes": []}
    needs_readback = owner.post(f"/documents/{document['id']}/records", expected=409, json=document_body).json()
    assert needs_readback["error"]["code"] == "capture_checks_changed"
    readback = needs_readback["error"]["details"]
    document_body.update(check_token=readback["check_token"], acknowledged_issue_codes=readback["requires_acknowledgement"])
    retry = {"Idempotency-Key": "sale-document-" + uuid4().hex}
    sale = owner.post(f"/documents/{document['id']}/records", expected=201, headers=retry, json=document_body).json()
    assert owner.post(f"/documents/{document['id']}/records", expected=201, headers=retry, json=document_body).json() == sale
    owner.post(f"/documents/{document['id']}/records", expected=409, headers=retry, json={**document_body, "data": {**raw, "amount": "999,00"}})
    assert sale["data"]["review_status"] == "proposed"
    assert {key: sale["data"][key] for key in ("date", "buyer", "unit", "currency")} == {"date": "2026-06-11", "buyer": "Fictional Mill", "unit": "kg", "currency": "EUR"}
    assert Decimal(str(sale["data"]["quantity"])) == Decimal("25.50")
    assert Decimal(str(sale["data"]["amount"])) == Decimal("125.50")
    assert sale["data"]["document_id"] == document["id"]
    assert "decimal_comma" in sale["data"]["document_review"]["acknowledged_issue_codes"]
    assert owner.get(f"/documents/{document['id']}/file").content == original
    assert owner.get("/facts").json()["total"] == 0, "Saving an invoice must not invent turnover, tax or stock facts."
    updated = owner.patch(f"/sales/{sale['id']}", json={"expected_revision": sale["revision"], "data": {"notes": "Printed total includes no inferred tax calculation."}}).json()
    owner.patch(f"/sales/{sale['id']}", expected=409, json={"expected_revision": sale["revision"], "data": {"amount": 999.00}})

    area = confirm(owner, owner.create("facts", fact_data(farm["id"], basis="static", period=None)))
    derived = confirm(owner, owner.create("facts", fact_data(farm["id"], key="printed_sale_total", value="125.50", unit="EUR", basis="calendar_date", period="2026-06-11", sources=[{
        "record_id": updated["id"], "record_kind": "sales", "revision": updated["revision"], "origin": "Manually read printed total, not quantity multiplied by an assumed price.",
    }])))
    report_body = {"farm_id": farm["id"], "pack_id": "farm-evidence", "reporting_year": 2026, "fact_ids": [area["id"], derived["id"]]}
    unreviewed = owner.post("/reports", expected=201, json=report_body).json()
    blocked = owner.post(f"/reports/{unreviewed['id']}/approve", expected=409, json={"expected_revision": unreviewed["revision"]}).json()
    assert "unreviewed_source_record" in json.dumps(blocked)
    sale = owner.post(f"/sales/{sale['id']}/review", json={"expected_revision": updated["revision"], "decision": "confirmed", "note": "Checked actual date, buyer, printed 125.50 EUR total and 25.50 kg against the authored original; VAT remains unknown."}).json()
    owner.post(f"/reports/{unreviewed['id']}/approve", expected=409, json={"expected_revision": unreviewed["revision"]})
    owner.patch(f"/sales/{sale['id']}", expected=409, json={"expected_revision": sale["revision"], "data": {"amount": 120.00}})
    # A source confirmation changes its revision. Refresh the citation explicitly.
    derived = confirm(owner, owner.create("facts", {**fact_data(farm["id"], key="printed_sale_total", value="125.50", unit="EUR", basis="calendar_date", period="2026-06-11", sources=[{
        "record_id": sale["id"], "record_kind": "sales", "revision": sale["revision"],
    }]), "supersedes_id": derived["id"]}))
    report_body["fact_ids"] = [area["id"], derived["id"]]

    # A credit is a separate dated financial entry, not negative stock movement.
    credit_raw = {**raw, "date": "2026-06-12", "title": "Separate authored sales credit note", "invoice_number": "AUTHORED-CREDIT-001", "amount": "-12,50", "quantity": None, "unit": None}
    negative = owner.post("/workflow/check", json={"kind": "sales", "data": credit_raw}).json()
    assert "negative_sale" in negative["requires_acknowledgement"]
    credit_raw["credit"] = True
    credit_check = owner.post("/workflow/check", json={"kind": "sales", "data": credit_raw}).json()
    credit_body = {"kind": "sales", "data": credit_raw, "check_token": credit_check["check_token"], "acknowledged_codes": credit_check["requires_acknowledgement"]}
    credit_retry = {"Idempotency-Key": "sale-credit-" + uuid4().hex}
    credit = owner.post("/workflow/capture", expected=201, headers=credit_retry, json=credit_body).json()["record"]
    assert owner.post("/workflow/capture", expected=201, headers=credit_retry, json=credit_body).json()["record"]["id"] == credit["id"]
    assert Decimal(str(credit["data"]["amount"])) == Decimal("-12.50") and credit["data"]["credit"] is True

    # Explicit column choices turn a CSV into proposed sales, never purchases.
    source_csv = b'buyer,date,product,quantity,unit,total,currency,invoice\r\nFictional Bakery,2026-06-13,Wheat,"2,00",kg,"20,00",EUR,AUTHORED-CSV-001\r\n'
    imported = owner.post("/imports", expected=201, files={"file": ("authored-outgoing.csv", source_csv, "text/csv")}, data={"farm_id": farm["id"], "grain": "One explicitly dated outgoing crop sale"}).json()
    assert owner.get("/sales").json()["total"] == 2
    mapping = owner.post(f"/imports/{imported['id']}/mapping", json={"expected_revision": imported["revision"], "target": "sales", "defaults": {"farm_id": farm["id"], "category": "crop", "credit": False}, "columns": [
        {"source_index": index, "target_field": field, "value_type": "number" if field in {"quantity", "amount"} else "date" if field == "date" else "text", **({"decimal_separator": ","} if field in {"quantity", "amount"} else {"date_format": "iso"} if field == "date" else {})}
        for index, field in enumerate(("buyer", "date", "title", "quantity", "unit", "amount", "currency", "invoice_number"))]}).json()
    preview = owner.get(f"/imports/{imported['id']}/preview").json()["items"][0]
    assert preview["target"] == "sales" and preview["can_save"] is True
    assert preview["row"]["data"]["original_row"][3:6] == ["2,00", "kg", "20,00"]
    promoted = owner.post(f"/imports/{imported['id']}/rows/review", json={"expected_revision": mapping["revision"], "note": "Selected outgoing sales and comma-decimal columns; actual source date is not an annual income summary.", "rows": [{"row_number": 1, "expected_revision": preview["row"]["revision"], "preview_token": preview["preview_token"], "acknowledged_issues": preview["requires_acknowledgement"]}]}).json()
    assert promoted["saved_count"] == 1 and promoted["failed_count"] == 0
    csv_sale = promoted["items"][0]["record"]
    assert csv_sale["data"]["review_status"] == "proposed"
    assert csv_sale["data"]["buyer"] == "Fictional Bakery" and csv_sale["data"]["unit"] == "kg"
    assert csv_sale["data"]["period"] == "2026-06-13" and Decimal(str(csv_sale["data"]["amount"])) == Decimal("20.00")
    assert any(source.get("import_id") == imported["id"] and source.get("row_number") == 1 for source in csv_sale["data"]["sources"])

    hidden = owner.create("sales", {**raw, "farm_id": private["id"], "quantity": 25.50, "amount": 125.50})
    outsider.get(f"/sales/{sale['id']}", expected=404)
    outsider.get(f"/documents/{document['id']}/file", expected=404)
    outsider.post(f"/sales/{credit['id']}/review", expected=404, json={"expected_revision": credit["revision"], "decision": "confirmed", "note": "Must not review another workspace."})
    invitation = owner.post("/invitations", expected=201, json={"email": adviser.email, "role": "adviser", "farm_ids": [farm["id"]], "permissions": ["read", "review"], "expires_in_days": 1}).json()
    adviser.post("/invitations/accept", json={"token": invitation["token"]})
    selected = {"X-Workspace-ID": owner.workspace_id}
    assert {item["id"] for item in adviser.get("/sales", headers=selected).json()["items"]} == {sale["id"], credit["id"], csv_sale["id"]}
    adviser.get(f"/sales/{hidden['id']}", headers=selected, expected=404)
    adviser.post("/sales", headers=selected, expected=404, json={"data": {**raw, "farm_id": private["id"], "quantity": 25.50, "amount": 125.50}})
    for record in (credit, csv_sale):
        reviewed = adviser.post(f"/sales/{record['id']}/review", headers=selected, json={"expected_revision": record["revision"], "decision": "confirmed", "note": "Checked this individual authored sale or credit and its source, not a calculated turnover or animal event."}).json()
        assert reviewed["data"]["reviewed_by"] == adviser.user_id
    owner.post(f"/invitations/{invitation['invitation']['id']}/revoke")
    assert adviser.client.get(f"/sales/{sale['id']}", headers=selected).status_code in (401, 403, 404)
    assert owner.get("/purchases").json()["total"] == 0 and owner.get("/operations").json()["total"] == 0
    assert all(item["data"].get("tax_amount") is None and item["data"].get("tax_rate") is None for item in owner.get("/sales").json()["items"])

    draft = owner.post("/reports", expected=201, json=report_body).json()
    issued = owner.post(f"/reports/{draft['id']}/approve", json={"expected_revision": draft["revision"]}).json()
    retained = {format: owner.get(f"/reports/{issued['id']}/download", params={"format": format}).content for format in ("json", "csv", "xlsx", "pdf")}
    snapshot = json.loads(retained["json"])["snapshot"]
    selected_fact = next(item for item in snapshot["facts"] if item["id"] == derived["id"])
    assert any(source.get("record_id") == sale["id"] and source.get("record_kind") == "sales" and source.get("revision") == sale["revision"] for source in selected_fact["data"]["sources"])
    assert next(field for field in snapshot["fields"] if field["fact_id"] == derived["id"])["value"] == "125.50"
    assert json.loads(retained["json"])["filing_status"] == "not_filed"
    rows = list(csv.DictReader(StringIO(retained["csv"].decode("utf-8-sig"))))
    assert any(row["fact_id"] == derived["id"] and json.loads(row["value_json"]) == "125.50" and json.loads(row["unit_json"]) == "EUR" for row in rows)
    workbook = load_workbook(BytesIO(retained["xlsx"]), read_only=True)
    try:
        assert any(derived["id"] in row and '"125.50"' in row for row in workbook.active.iter_rows(values_only=True))
    finally:
        workbook.close()
    assert "125.50" in "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(retained["pdf"])).pages)

    correction_data = {**raw, "amount": "120.00", "quantity": "25.50", "notes": "A separately checked corrected invoice, not an overwrite.", "sources": [{"kind": "correction", "record_id": sale["id"], "record_kind": "sales", "revision": sale["revision"]}]}
    correction_check = owner.post("/workflow/check", json={"kind": "sales", "data": correction_data}).json()
    correction = owner.post("/workflow/capture", expected=201, json={"kind": "sales", "data": correction_data, "check_token": correction_check["check_token"], "acknowledged_codes": correction_check["requires_acknowledgement"]}).json()["record"]
    corrected = owner.post(f"/sales/{correction['id']}/review", json={"expected_revision": correction["revision"], "decision": "confirmed", "note": "Read corrected printed total 120.00 EUR; original reviewed invoice retained."}).json()
    owner.post(f"/sales/{sale['id']}/review", json={"expected_revision": sale["revision"], "decision": "rejected", "note": "Explicitly replaced by the separately reviewed correction; retain old source and issued output."})
    assert issued["id"] in {item["id"] for item in owner.get(f"/sales/{sale['id']}/reports").json()["items"]}
    assert any(item["data"]["review_status"] == "confirmed" and Decimal(str(item["data"]["amount"])) == Decimal("125.50") for item in owner.get(f"/sales/{sale['id']}/history").json()["items"])
    corrected_fact = confirm(owner, owner.create("facts", {**fact_data(farm["id"], key="printed_sale_total", value="120.00", unit="EUR", basis="calendar_date", period="2026-06-11", sources=[{"record_id": corrected["id"], "record_kind": "sales", "revision": corrected["revision"]}]), "supersedes_id": derived["id"]}))
    amendment = owner.post(f"/reports/{issued['id']}/reissue", expected=201, json={"expected_revision": issued["revision"], "fact_ids": [area["id"], corrected_fact["id"]]}).json()
    assert amendment["id"] != issued["id"] and amendment["data"]["status"] == "draft"
    owner.post(f"/reports/{amendment['id']}/approve", json={"expected_revision": amendment["revision"]})
    for format in ("json", "csv", "xlsx"):
        assert owner.get(f"/reports/{issued['id']}/download", params={"format": format}).content == retained[format]
    old_pdf = owner.get(f"/reports/{issued['id']}/download", params={"format": "pdf"}).content
    assert [page.extract_text() for page in PdfReader(BytesIO(old_pdf)).pages] == [page.extract_text() for page in PdfReader(BytesIO(retained["pdf"])).pages]
    assert owner.get(f"/sales/{sale['id']}").json()["data"]["amount"] == sale["data"]["amount"]
    outsider.get(f"/reports/{issued['id']}/download", params={"format": "json"}, expected=404)
