"""Scope cases 1, 4, 5, 9: reviewed outputs and retained correction history."""

import csv
from io import BytesIO, StringIO
import json
from uuid import uuid4

import httpx
from openpyxl import load_workbook
from pypdf import PdfReader

from conftest import confirm, fact_data, farm_data


def test_review_two_outputs_correct_and_preserve_issued_exports(farmers, service_url):
    owner = farmers("Report owner")
    recipient = farmers("Named recipient")
    farm = owner.create("farms", farm_data())
    private_request_note = "PRIVATE-REQUEST-NOTE-" + uuid4().hex
    private_source_name = "UNSELECTED-SOURCE-" + uuid4().hex + ".txt"
    source_document = owner.post("/documents", files={"file": (private_source_name, b"Authored unselected private evidence.\n", "text/plain")}, data={"farm_id": farm["id"]}, expected=201).json()
    # These country-neutral fixture packs make no national form or legal claim.
    for pack_id in ("farm-evidence", "inspection-evidence"):
        pack = owner.get(f"/packs/{pack_id}").json()
        assert pack["country"] is None
        assert pack["capability"] == "generic_dossier"
    request = owner.create("requests", {
        "farm_id": farm["id"], "title": "Authored annual area evidence request",
        "recipient": "Fictional review office", "country": "DE", "region": None,
        "reporting_year": 2026, "deadline": "2026-12-01", "pack_id": "farm-evidence",
        "requirements": [{"key": "area_ha", "label": "Recorded holding area", "required": True}],
        "status": "draft", "notes": private_request_note,
    })
    proposed = owner.create("facts", fact_data(farm["id"], value="3.75", basis="static", period=None, sources=[{
        "document_id": source_document["id"], "file_name": private_source_name, "row": 7,
        "origin": "Unselected private source must not accompany a shared output.",
    }]))
    report_body = {"farm_id": farm["id"], "request_id": request["id"], "pack_id": "farm-evidence", "reporting_year": 2026, "fact_ids": [proposed["id"]]}
    unreviewed = owner.post("/reports", json=report_body, expected=201).json()
    owner.post(f"/reports/{unreviewed['id']}/approve", json={"expected_revision": unreviewed["revision"]}, expected=409)
    confirmed = confirm(owner, proposed)
    # Explicit review changes the selected revision; the old draft stays blocked.
    owner.post(f"/reports/{unreviewed['id']}/approve", json={"expected_revision": unreviewed["revision"]}, expected=409)

    context_preview = owner.post("/reports", json={**report_body, "fact_ids": [confirmed["id"]]}, expected=201).json()
    request = owner.patch(f"/requests/{request['id']}", json={"expected_revision": request["revision"], "data": {
        "requirements": [{"key": "area_ha", "label": "Recorded holding area", "required": True, "unit": "m2", "temporal_basis": "campaign_year"}],
    }}).json()
    owner.post(f"/reports/{context_preview['id']}/approve", json={"expected_revision": context_preview["revision"]}, expected=409)
    wrong_requirement = owner.post("/reports", json={**report_body, "fact_ids": [confirmed["id"]]}, expected=201).json()
    owner.post(f"/reports/{wrong_requirement['id']}/approve", json={"expected_revision": wrong_requirement["revision"]}, expected=409)
    assert owner.get(f"/facts/{confirmed['id']}").json()["data"]["unit"] == "ha"
    request = owner.patch(f"/requests/{request['id']}", json={"expected_revision": request["revision"], "data": {
        "requirements": [{"key": "area_ha", "label": "Recorded holding area", "required": True, "unit": "ha", "temporal_basis": "static"}],
    }}).json()

    issued = []
    for pack_id in ("farm-evidence", "inspection-evidence"):
        body = {**report_body, "pack_id": pack_id, "fact_ids": [confirmed["id"]]}
        if pack_id == "inspection-evidence":
            body.pop("request_id")
        draft = owner.post("/reports", json=body, expected=201).json()
        report = owner.post(f"/reports/{draft['id']}/approve", json={"expected_revision": draft["revision"]}).json()
        assert report["data"]["status"] == "issued"
        assert report["data"]["filing_status"] == "not_filed"
        assert report["data"]["approval"]["actor_id"] == owner.user_id
        same = owner.post(f"/reports/{report['id']}/approve", json={"expected_revision": report["revision"]}).json()
        assert same == report
        issued.append(report)
    assert owner.get(f"/requests/{request['id']}").json()["data"]["status"] == "ready"

    retained = {}
    for report in issued:
        downloads = {format: owner.get(f"/reports/{report['id']}/download", params={"format": format}) for format in ("json", "csv", "xlsx", "pdf")}
        exported = downloads["json"].json()
        assert exported["snapshot"] == report["data"]["snapshot"]
        assert exported["filing_status"] == "not_filed"
        field = next(item for item in exported["snapshot"]["fields"] if item["fact_id"] == confirmed["id"])
        assert field["value"] == "3.75" and field["unit"] == "ha" and field["period"] is None
        assert field["temporal_basis"] == "static"
        assert field["fact_revision"] == confirmed["revision"]
        rows = list(csv.DictReader(StringIO(downloads["csv"].content.decode("utf-8-sig"))))
        assert any(json.loads(row["value_json"]) == "3.75" and row["fact_id"] == confirmed["id"] and json.loads(row["unit_json"]) == "ha" for row in rows)
        workbook = load_workbook(BytesIO(downloads["xlsx"].content), read_only=True, data_only=False)
        try:
            cells = list(workbook.active.iter_rows())
            header = [cell.value for cell in cells[0]]
            sheet_rows = [dict(zip(header, [cell.value for cell in row])) for row in cells[1:]]
            assert any(json.loads(row["value_json"]) == "3.75" and row["fact_id"] == confirmed["id"] for row in sheet_rows)
            assert all(cell.data_type != "f" for row in cells for cell in row)
        finally:
            workbook.close()
        pdf_text = "\n".join(page.extract_text() or "" for page in PdfReader(BytesIO(downloads["pdf"].content)).pages)
        assert "3.75" in pdf_text and "ha" in pdf_text
        assert "Report reference " + report["id"][:8] in pdf_text
        assert "acceptance" in pdf_text.lower()
        for format, response in downloads.items():
            assert response.headers["x-report-snapshot-sha256"] == report["data"]["checksum"]
            assert "no-store" in response.headers["cache-control"].lower()
            recipient.get(f"/reports/{report['id']}/download", params={"format": format}, expected=404)
        retained[report["id"]] = {format: response.content for format, response in downloads.items()}

    correction_data = {**fact_data(farm["id"], value="4.25", basis="static", period=None), "supersedes_id": confirmed["id"]}
    correction = owner.create("facts", correction_data)
    corrected = confirm(owner, correction, "A new survey found 4.25 ha; replaces the earlier measured declaration.")
    dependencies = owner.get(f"/facts/{confirmed['id']}/reports").json()["items"]
    assert {report["id"] for report in issued}.issubset({report["id"] for report in dependencies})
    history = owner.get(f"/facts/{confirmed['id']}/history").json()["items"]
    assert any(item["data"]["review_status"] == "confirmed" and item["data"]["value"] == "3.75" for item in history)

    for report in issued:
        old = owner.get(f"/reports/{report['id']}").json()
        assert old["data"]["snapshot"] == report["data"]["snapshot"]
        assert old["data"]["checksum"] == report["data"]["checksum"]
        for format in ("json", "csv", "xlsx"):
            assert owner.get(f"/reports/{report['id']}/download", params={"format": format}).content == retained[report["id"]][format]
        # PDF may contain generation timestamps, so compare its extracted content.
        old_pdf = owner.get(f"/reports/{report['id']}/download", params={"format": "pdf"}).content
        assert [page.extract_text() for page in PdfReader(BytesIO(old_pdf)).pages] == [page.extract_text() for page in PdfReader(BytesIO(retained[report["id"]]["pdf"])).pages]
        revised = owner.post(f"/reports/{report['id']}/reissue", json={"expected_revision": old["revision"], "fact_ids": [corrected["id"]]}, expected=201).json()
        assert revised["id"] != report["id"]
        assert revised["data"]["correction_of"] == report["id"]
        assert revised["data"]["status"] == "draft"
        assert any(field["value"] == "4.25" for field in revised["data"]["snapshot"]["fields"])
        owner.post(f"/reports/{revised['id']}/approve", json={"expected_revision": revised["revision"]})

    first = issued[0]
    receipt_body = {"expected_revision": first["revision"], "reference": "FICTIONAL-MANUAL-REFERENCE-2026", "date": "2026-10-16"}
    receipt = owner.post(f"/reports/{first['id']}/receipts", json=receipt_body, expected=201).json()
    assert receipt["data"]["authority_acceptance_verified"] is False
    assert receipt["data"]["status"] == "manually_recorded"
    retry = owner.post(f"/reports/{first['id']}/receipts", json=receipt_body, expected=201).json()
    assert retry["id"] == receipt["id"]
    assert owner.get(f"/requests/{request['id']}").json()["data"]["status"] == "responded"

    share = owner.post("/shares", json={"report_id": first["id"], "recipient_email": recipient.email, "expires_in_days": 1}, expected=201).json()
    link = "/shares/access/" + share["token"]
    shared = recipient.get(link)
    assert shared.json()["report"]["id"] == first["id"]
    shared_download = recipient.get(link + "/download", params={"format": "json"})
    assert shared_download.json()["snapshot"] == first["data"]["snapshot"]
    for exposed in (shared.text, shared_download.text):
        assert private_request_note not in exposed
        assert private_source_name not in exposed
        assert source_document["id"] not in exposed
        assert "Unselected private source must not accompany a shared output." not in exposed
    with httpx.Client(base_url=service_url + "/", timeout=30) as anonymous:
        assert anonymous.get(link.lstrip("/")).status_code == 401
    owner.post(f"/shares/{share['share']['id']}/revoke")
    assert recipient.client.get(link).status_code in (403, 404)
    assert recipient.client.get(link + "/download", params={"format": "json"}).status_code in (403, 404)
    assert owner.get(f"/reports/{first['id']}").json()["data"]["snapshot"] == first["data"]["snapshot"]
    other_document = owner.post("/documents", files={"file": ("unshared-private-note.txt", b"This authored original was not shared.\n", "text/plain")}, data={"farm_id": farm["id"]}, expected=201).json()
    originals = owner.post("/shares", expected=201, json={"document_ids": [source_document["id"]], "recipient_email": recipient.email, "expires_in_days": 1}).json()
    originals_link = "/shares/access/" + originals["token"]
    selected_original = recipient.get(originals_link + f"/documents/{source_document['id']}/download")
    assert selected_original.content == b"Authored unselected private evidence.\n"
    recipient.get(originals_link + f"/documents/{other_document['id']}/download", expected=404)
    owner.post(f"/shares/{originals['share']['id']}/revoke")
    assert recipient.client.get(originals_link + f"/documents/{source_document['id']}/download").status_code in (403, 404)
