"""Authored Czech journal acceptance through real HTTP and tenant-scoped storage.

Oracles come from the retained UKZUZ record example (publication 2026-03-30):
one actual application per entry, conditional locations, two/three-decimal source
precision, optional total quantity, and pending efficacy distinct from recorded 0.
This fixture is fictional and never asserts a product's current authorisation.
"""

from copy import deepcopy
from decimal import Decimal
from hashlib import sha256
from io import BytesIO
import json
import os
import subprocess
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from conftest import PROJECT, confirm, fact_data, farm_data, human_browser


PACK = "cz-ppp-journal-2026"
SOURCE_HASH = "bae53ed12a7b06029b77ceb9dc216d498933d79d6f0f511404bc04ca3a461403"


def application(farm_id, location="lpis_land", *, title="Fictional first treatment"):
    details = {
        "location_kind": location, "cadastral_area": "Authored cadastral area 001",
        "subject_kind": "crop", "sown_area_ha": 20.35, "sowing_date": "2025-08-10",
        "crop_end": {"status": "pending"}, "target_organism_or_purpose": "Fictional recorded target",
        "mixture": "S", "efficacy": {"status": "pending"},
        "original_numbers": {"dose": "0.500", "treated_size": "20.35", "sown_area_ha": "20.35"},
    }
    if location == "lpis_land":
        details.update(dpb_square="000012", dpb_code="0007/01")
    elif location == "other_land":
        details["cadastral_parcel_numbers"] = ["0008/02", "0009"]
    elif location == "forest":
        details.update(subject_kind="forest", lhc_code="000055", jprl={
            "oddeleni": "001", "dilec": "A", "porost": "02", "porostni_skupina": "003"},
            crop_end={"status": "not_applicable"})
        details.pop("sowing_date")
    elif location == "object":
        details.update(subject_kind="other", cadastral_parcel_numbers=["0010/03"],
                       object_type="sklad", object_description="Fictional north store, bay 01")
        details.pop("sown_area_ha")
        details.pop("sowing_date")
        details["original_numbers"].pop("sown_area_ha")
    data = {
        "farm_id": farm_id, "title": title, "type": "pesticide_application",
        "date": "2026-06-15", "period": "2026-06-15", "temporal_basis": "calendar_date",
        "product": "Fictional product — not an authorisation claim", "product_authorisation": "AUTHORED-ONLY",
        "crop": "Fictional crop", "dose": 0.5, "dose_unit": "L/ha", "treated_size": 20.35,
        "treated_unit": "ha", "quantity": None, "unit": None,
        "treatment_type": "surface", "treated_location": "Fictional exact treatment location",
        "label_time_required": False, "label_growth_stage_required": False,
        "sources": [{"origin": "Authored dated field notebook; not an invoice or official submission"}],
        "cz_ppp": details,
    }
    if location == "object":
        data.update(treatment_type="closed_space", dose_unit="L/m3", treated_unit="m3")
    return data


def capture(actor, data):
    body = {"kind": "operations", "data": data}
    check = actor.post("/workflow/check", json=body).json()
    assert check["can_save"], check["issues"]
    payload = {**body, "check_token": check["check_token"], "acknowledged_codes": check["requires_acknowledgement"]}
    retry = {"Idempotency-Key": "cz-capture-" + uuid4().hex}
    result = actor.post("/workflow/capture", expected=201, headers=retry, json=payload).json()
    assert actor.post("/workflow/capture", expected=201, headers=retry, json=payload).json() == result
    record = result["record"]
    assert record["data"]["review_status"] == "proposed"
    return record, check


def review_operation(actor, record, expected=200):
    return actor.post(f"/operations/{record['id']}/review", expected=expected, json={
        "expected_revision": record["revision"], "decision": "confirmed",
        "note": "Compared this actual application with the exact original identifiers, units and pending follow-ups.",
    }).json()


def test_czech_repeated_applications_pending_followup_corrections_and_immutable_journal(farmers, service_url):
    owner, outsider, adviser = (farmers(name) for name in ("Czech journal owner", "Unrelated reader", "Read-only farm adviser"))
    farm = owner.create("farms", {**farm_data("Fictional Czech holding"), "country": "CZ", "region": None})
    other = owner.create("farms", {**farm_data("Different Czech holding"), "country": "CZ", "region": None})
    catalogue = owner.get("/countries/CZ").json()
    input_workflow = next(item for item in catalogue["modules"] if item["workflow_id"] == "inputs")
    assert input_workflow["mapping_status"] == "partial_mapping" and PACK in input_workflow["pack_ids"]
    assert any(source.get("sha256") == SOURCE_HASH for source in input_workflow["official_links"])
    request = owner.create("requests", {"farm_id": farm["id"], "title": "Authored repeated Czech journal",
        "country": "CZ", "region": None, "reporting_year": 2026, "recipient": "Fictional private reviewer",
        "pack_id": PACK, "requirements": [], "deadline": None, "status": "draft"})
    invitation = owner.post("/invitations", expected=201, json={"email": adviser.email, "role": "adviser",
        "farm_ids": [farm["id"]], "permissions": ["read"], "expires_in_days": 1}).json()
    adviser.post("/invitations/accept", json={"token": invitation["token"]})
    shared_scope = {"X-Workspace-ID": owner.workspace_id}
    records = []
    raw = []
    for index, location in enumerate(("lpis_land", "lpis_land", "other_land", "forest", "object")):
        data = application(farm["id"], location, title=f"Fictional actual application {index + 1}")
        if index == 1:
            data["cz_ppp"]["mixture"] = "TM"
            data["cz_ppp"]["efficacy"] = {"status": "recorded", "value": 0}
            data["cz_ppp"]["crop_end"] = {"status": "recorded", "date": "2026-09-15"}
        proposed, check = capture(owner, data)
        if location == "object":
            assert any(issue["code"] == "cz_ppp_pending_crop_end" for issue in check["issues"])
        assert proposed["data"]["quantity"] is None and proposed["data"]["unit"] is None
        assert proposed["data"]["cz_ppp"] == data["cz_ppp"]
        assert Decimal(str(proposed["data"]["dose"])) == Decimal("0.500")
        records.append(review_operation(owner, proposed)); raw.append(data)
    assert len({item["id"] for item in records}) == 5
    assert records[0]["data"]["date"] == records[1]["data"]["date"]
    assert records[0]["data"]["cz_ppp"]["dpb_code"] == records[1]["data"]["cz_ppp"]["dpb_code"]
    outsider.get(f"/operations/{records[0]['id']}", expected=404)
    adviser.get(f"/operations/{records[0]['id']}", headers=shared_scope)
    adviser.get(f"/farms/{other['id']}", headers=shared_scope, expected=404)

    # Malformed/overprecise values must not be rounded, converted or treated as pending.
    for change in (
        {"efficacy": {"status": "pending", "value": 0}},
        {"efficacy": {"status": "recorded", "value": 2}},
        {"efficacy": {"status": "recorded", "value": False}},
        {"crop_end": {"status": "recorded", "date": "not-an-actual-date"}},
        {"original_numbers": {"dose": "0.5001", "treated_size": "20.35", "sown_area_ha": "20.35"}},
    ):
        invalid = deepcopy(raw[0]); invalid["cz_ppp"].update(change)
        checked = owner.post("/workflow/check", json={"kind": "operations", "data": invalid}).json()
        assert checked["can_save"] is False, "Impossible national details must not silently become usable evidence."
    invalid = deepcopy(raw[0]); invalid.pop("cz_ppp")
    invalid.update(date="not-an-actual-date", period="not-an-actual-date")
    invalid_body = {"kind": "operations", "data": invalid}
    checked = owner.post("/workflow/check", json=invalid_body).json()
    assert checked["can_save"] is False, "Canonical actual dates remain validated without national details."
    owner.post("/workflow/capture", expected=422, json={**invalid_body,
        "check_token": checked["check_token"], "acknowledged_codes": checked["requires_acknowledgement"]})
    changed_details = deepcopy(raw[0]["cz_ppp"])
    changed_details["original_numbers"]["dose"] = "0.600"
    owner.patch(f"/operations/{records[0]['id']}", expected=409,
                json={"expected_revision": records[0]["revision"], "data": {"dose": 0.6, "cz_ppp": changed_details}})

    compact = {"view": "summary"}
    report_body = {"farm_id": farm["id"], "request_id": request["id"], "pack_id": PACK,
                   "reporting_year": 2026, "operation_ids": [item["id"] for item in records]}
    no_holder = owner.post("/reports", expected=201, params=compact, json=report_body).json()
    missing = owner.get(f"/reports/{no_holder['id']}?view=review").json()
    assert missing["data"]["can_approve"] is False, "Farm display name must not supply an invented holder."
    owner.post(f"/reports/{no_holder['id']}/approve", expected=409, json={"expected_revision": no_holder["revision"]})
    holder_value = {"kind": "legal_person", "name": "Fictional Czech grower s.r.o.", "address": "Authored business address 01",
                    "identification_number": {"status": "provided", "value": "00001234"}}
    holder = confirm(owner, owner.create("facts", fact_data(farm["id"], key="cz.ppp.holder", value=holder_value,
                                                         unit=None, basis="static", period=None)))
    # A small view must page the whole authorized ledger, not silently omit entries.
    paged_entries = []
    for offset, next_offset, size in ((0, 2, 2), (2, 4, 2), (4, None, 1)):
        guided = owner.get("/workflow/questions", params={"request_id": request["id"],
            "journal_limit": 2, "journal_offset": offset}).json()
        task = next(item for item in guided["journals"] if item["pack_id"] == PACK)
        assert task["pagination"] == {"total": 5, "limit": 2, "offset": offset, "next_offset": next_offset}
        assert task["counts"]["total"] == 5 and task["counts"]["confirmed"] == 5
        assert len(task["entries"]) == size
        paged_entries.extend(task["entries"])
    assert [entry["operation_id"] for entry in paged_entries] == [item["id"] for item in sorted(
        records, key=lambda item: (item["data"]["date"], item["id"]))]
    assert len({entry["operation_id"] for entry in paged_entries}) == 5
    # Each official location family has its own required identity, rather than
    # treating an arbitrary cadastral label as an LPIS/forest/object identifier.
    incomplete = []
    for location, missing_field in (("lpis_land", "dpb_code"), ("other_land", "cadastral_parcel_numbers"),
                                    ("forest", "lhc_code"), ("object", "object_description")):
        data = application(farm["id"], location, title="Incomplete location for later source completion")
        data["cz_ppp"].pop(missing_field)
        proposed, check = capture(owner, data)
        assert any(issue.get("field") == "cz_ppp." + missing_field for issue in check["issues"])
        incomplete.append(proposed)
    guided = owner.get("/workflow/questions", params={"request_id": request["id"]}).json()
    journal_task = next(item for item in guided["journals"] if item["pack_id"] == PACK)
    assert journal_task["source"]["sha256"] == SOURCE_HASH
    guided_entries = {entry["operation_id"]: entry for entry in journal_task["entries"]}
    assert {item["id"] for item in records + incomplete} <= set(guided_entries)
    assert all(guided_entries[item["id"]]["issues"] for item in incomplete)
    for proposed in incomplete:
        reviewed = review_operation(owner, proposed)
        blocked = owner.post("/reports", expected=201, params=compact,
                             json={**report_body, "operation_ids": [reviewed["id"]]}).json()
        assert not owner.get(f"/reports/{blocked['id']}?view=review").json()["data"]["can_approve"]
        owner.post(f"/reports/{blocked['id']}/approve", expected=409,
                   json={"expected_revision": blocked["revision"]})
    selected = ",".join(item["id"] for item in records)
    preflight = owner.get(f"/requests/{request['id']}/checks", params={"selected_operation_ids": selected}).json()
    assert set(preflight["selection"]["operation_ids"]) == {item["id"] for item in records}
    wrong_farm, _ = capture(owner, application(other["id"]))
    wrong_farm = review_operation(owner, wrong_farm)
    owner.post("/reports", expected=422, json={**report_body, "operation_ids": [wrong_farm["id"]]})
    old_year = application(farm["id"]); old_year.update(date="2025-06-15", period="2025-06-15")
    old_year["cz_ppp"]["sowing_date"] = "2024-08-10"
    past, _ = capture(owner, old_year); past = review_operation(owner, past)
    owner.post("/reports", expected=422, params=compact,
               json={**report_body, "operation_ids": [past["id"]]})

    draft = owner.post("/reports", expected=201, params=compact, json=report_body).json()
    metadata = owner.get(f"/reports/{draft['id']}?view=review").json()
    binding = metadata["data"]["review"]
    assert metadata["data"]["can_approve"], "Legitimate pending efficacy/end dates alone must not block a manual journal."
    outsider.get(f"/reports/{draft['id']}/review", expected=404, params={"section": "operations", "expected_revision": draft["revision"],
        "snapshot_checksum": binding["snapshot_checksum"], "check_fingerprint": binding["check_fingerprint"]})
    params = {"section": "operations", "expected_revision": draft["revision"], "snapshot_checksum": binding["snapshot_checksum"],
              "check_fingerprint": binding["check_fingerprint"], "limit": 5, "offset": 0}
    page = owner.get(f"/reports/{draft['id']}/review", params=params)
    assert len(page.content) <= 2097152
    items = page.json()["items"]
    assert {item["id"] for item in items} == {item["id"] for item in records}
    for item in items:
        context = item["journal_entry"]
        assert context["source"]["sha256"] == SOURCE_HASH and context["columns"]
        assert context["entry"]["operation_id"] == item["id"]
        assert context["entry"]["operation_revision"] == item["revision"]
    approve = {"expected_revision": draft["revision"], "snapshot_checksum": binding["snapshot_checksum"], "check_fingerprint": binding["check_fingerprint"]}
    adviser.post(f"/reports/{draft['id']}/approve", headers=shared_scope, json=approve, expected=403)
    issued = owner.post(f"/reports/{draft['id']}/approve", json=approve).json()
    snapshot = issued["data"]["snapshot"]
    journal = snapshot["journals"][0]
    assert journal["source"]["sha256"] == SOURCE_HASH
    assert journal["holder"]["fact_id"] == holder["id"] and journal["holder"]["value"] == holder_value
    columns = {column["key"]: column for column in journal["columns"]}
    assert columns["dose"]["decimal_places"] == 3
    assert columns["treated_size"]["decimal_places"] == 2
    assert columns["sown_area_ha"]["unit"] == "ha"
    assert columns["quantity"]["optional"] is True
    entries = {entry["operation_id"]: entry for entry in journal["entries"]}
    assert set(entries) == {item["id"] for item in records}
    assert entries[records[0]["id"]]["data"]["cz_ppp"]["efficacy"] == {"status": "pending"}
    assert entries[records[1]["id"]]["data"]["cz_ppp"]["efficacy"] == {"status": "recorded", "value": 0}
    assert entries[records[1]["id"]]["data"]["cz_ppp"]["crop_end"]["date"] == "2026-09-15"
    assert all(entry["data"]["quantity"] is None for entry in entries.values())
    assert not set(entries) & {entry["id"] for entry in incomplete}, "Unselected incomplete applications must not enter the report."
    assert entries[records[0]["id"]]["values"]["dose"] == "0.500"
    assert entries[records[0]["id"]]["values"]["treated_size"] == "20.35"
    assert entries[records[0]["id"]]["original_numbers"]["dose"] == "0.500"
    assert entries[records[4]["id"]]["values"]["treated_unit"] == "m3"
    assert any(issue["code"] == "cz_ppp_pending_crop_end" for issue in entries[records[4]["id"]]["issues"])
    original_exports = {fmt: owner.get(f"/reports/{issued['id']}/download", params={"format": fmt}).content
                        for fmt in ("json", "csv", "pdf", "xlsx")}
    assert json.loads(original_exports["json"])["snapshot"]["journals"] == snapshot["journals"]
    assert "0.500" in original_exports["csv"].decode("utf-8-sig")
    workbook = load_workbook(BytesIO(original_exports["xlsx"]), read_only=True, data_only=True)
    assert "0.500" in " ".join(str(cell) for sheet in workbook for row in sheet.values for cell in row)
    assert "Fictional" in " ".join(page.extract_text() or "" for page in PdfReader(BytesIO(original_exports["pdf"])).pages)

    # Keeping a farm's manual journal does not require a letter or fabricated request.
    manual_body = {"farm_id": farm["id"], "pack_id": PACK, "reporting_year": 2026,
        "operation_ids": report_body["operation_ids"], "fact_ids": [holder["id"]]}
    manual = owner.post("/reports", expected=201, params=compact, json=manual_body).json()
    assert manual["data"]["request_id"] is None
    manual_review = owner.get(f"/reports/{manual['id']}?view=review").json()["data"]
    assert manual_review["can_approve"] and manual["data"]["status"] == "draft"
    manual_binding = manual_review["review"]
    manual_page = owner.get(f"/reports/{manual['id']}/review", params={"section": "operations", "offset": 0, "limit": 5,
        "expected_revision": manual["revision"], "snapshot_checksum": manual_binding["snapshot_checksum"],
        "check_fingerprint": manual_binding["check_fingerprint"]}).json()
    assert {item["id"] for item in manual_page["items"]} == set(manual_body["operation_ids"])
    manual_issued = owner.post(f"/reports/{manual['id']}/approve", json={"expected_revision": manual["revision"],
        "snapshot_checksum": manual_binding["snapshot_checksum"], "check_fingerprint": manual_binding["check_fingerprint"]}).json()
    assert manual_issued["data"]["request_id"] is None and manual_issued["data"]["status"] == "issued"
    assert manual_issued["data"]["snapshot"]["request"] is None

    # A correction remains a distinct proposal; confirmation, not capture, changes active evidence.
    original = records[0]
    pending_report = owner.post("/reports", expected=201, params=compact, json=report_body).json()
    pending_binding = owner.get(f"/reports/{pending_report['id']}?view=review").json()["data"]["review"]
    corrected = deepcopy(raw[0]); corrected.update(supersedes_id=original["id"], dose=0.6)
    corrected["cz_ppp"]["original_numbers"]["dose"] = "0.600"
    successor, _ = capture(owner, corrected)
    competing, _ = capture(owner, corrected)
    before_confirmation = owner.get(f"/operations/{original['id']}").json()
    assert before_confirmation["data"] == original["data"] and before_confirmation["revision"] == original["revision"]
    assert before_confirmation["active"] is True and before_confirmation["confirmed_successor_id"] is None
    assert not before_confirmation["data"].get("superseded_by")
    replacement = review_operation(owner, successor)
    after_confirmation = owner.get(f"/operations/{original['id']}").json()
    assert after_confirmation["data"] == original["data"] and after_confirmation["revision"] == original["revision"]
    assert after_confirmation["active"] is False and after_confirmation["confirmed_successor_id"] == replacement["id"]
    assert review_operation(owner, competing, expected=409)["error"]["code"] == "operation_correction_conflict"
    owner.post(f"/operations/{replacement['id']}/review", expected=409,
               json={"expected_revision": replacement["revision"] - 1, "decision": "confirmed", "note": "Stale review cannot overwrite a later revision."})
    assert owner.get(f"/reports/{pending_report['id']}?view=review").json()["data"]["can_approve"] is False
    owner.post(f"/reports/{pending_report['id']}/approve", expected=409, params=compact,
        json={"expected_revision": pending_report["revision"], "snapshot_checksum": pending_binding["snapshot_checksum"],
              "check_fingerprint": pending_binding["check_fingerprint"]})
    inactive = owner.post("/reports", expected=422, params=compact, json=report_body).json()
    assert inactive["error"]["code"] == "journal_active_entries_required"
    corrected_ids = [replacement["id"], *[item["id"] for item in records[1:]]]
    amendment = owner.post(f"/reports/{issued['id']}/reissue", expected=201, params=compact,
        json={"expected_revision": issued["revision"], "snapshot_checksum": issued["data"]["checksum"], "operation_ids": corrected_ids}).json()
    retained = owner.get(f"/reports/{amendment['id']}").json()["data"]["snapshot"]["journals"][0]
    assert {entry["operation_id"] for entry in retained["entries"]} == set(corrected_ids)
    manual_amendment = owner.post(f"/reports/{manual_issued['id']}/reissue", expected=201, params=compact,
        json={"expected_revision": manual_issued["revision"], "snapshot_checksum": manual_issued["data"]["checksum"],
              "operation_ids": corrected_ids, "fact_ids": [holder["id"]]}).json()
    assert manual_amendment["data"]["request_id"] is None
    assert manual_amendment["data"]["status"] == "draft" and manual_amendment["data"]["approval"] is None
    manual_amended_data = owner.get(f"/reports/{manual_amendment['id']}").json()["data"]
    assert manual_amended_data["snapshot"]["request"] is None
    assert {entry["operation_id"] for entry in manual_amended_data["snapshot"]["journals"][0]["entries"]} == set(corrected_ids)
    # The prior journal's review cannot approve its fresh correction draft.
    owner.post(f"/reports/{manual_amendment['id']}/approve", expected=409, json={
        "expected_revision": manual_amendment["revision"], "snapshot_checksum": manual_binding["snapshot_checksum"],
        "check_fingerprint": manual_binding["check_fingerprint"]})
    fresh_manual_review = owner.get(f"/reports/{manual_amendment['id']}?view=review").json()["data"]
    assert fresh_manual_review["can_approve"]
    fresh_binding = fresh_manual_review["review"]
    approved_manual_amendment = owner.post(f"/reports/{manual_amendment['id']}/approve", json={
        "expected_revision": manual_amendment["revision"], "snapshot_checksum": fresh_binding["snapshot_checksum"],
        "check_fingerprint": fresh_binding["check_fingerprint"]}).json()
    assert approved_manual_amendment["data"]["status"] == "issued" and approved_manual_amendment["data"]["request_id"] is None
    for fmt, content in original_exports.items():
        assert sha256(owner.get(f"/reports/{issued['id']}/download", params={"format": fmt}).content).digest() == sha256(content).digest()

    # Connected assistants use the same typed checks and cannot confirm evidence.
    assistant_data = application(farm["id"], title="Authored assistant-staged application")
    arguments = {"farm_id": farm["id"], "expected_farm_revision": farm["revision"],
                 "kind": "operations", "data": assistant_data}
    assistant_check = owner.post("/assistant/execute/check_capture", json={"arguments": arguments}).json()
    assert assistant_check["can_save"] and assistant_check["requires_acknowledgement"]
    assert assistant_check["normalized_data"]["cz_ppp"]["original_numbers"]["dose"] == "0.500"
    holder_arguments = {**arguments, "kind": "facts", "data": fact_data(farm["id"], key="cz.ppp.holder",
        value=holder_value, unit=None, basis="static", period=None)}
    assert owner.post("/assistant/execute/check_capture", json={"arguments": holder_arguments}).json()["can_save"]
    for malformed in (
        {**assistant_data, "cz_ppp": {**assistant_data["cz_ppp"], "arbitrary_nested": {"confirm": True}}},
        {**assistant_data, "cz_ppp": {**assistant_data["cz_ppp"], "efficacy": {"status": "pending", "unknown": "not allowed"}}},
    ):
        invalid = owner.post("/assistant/execute/check_capture", expected=422,
            json={"arguments": {**arguments, "data": malformed}}).json()
        assert invalid["error"]["code"] == "invalid_tool_arguments"
    wrong_kind = owner.post("/assistant/execute/check_capture", expected=422,
        json={"arguments": {**arguments, "kind": "purchases"}}).json()
    assert wrong_kind["error"]["code"] == "invalid_tool_arguments"
    invalid_business = deepcopy(assistant_data)
    invalid_business["cz_ppp"]["efficacy"] = {"status": "recorded", "value": 2}
    assert not owner.post("/assistant/execute/check_capture",
        json={"arguments": {**arguments, "data": invalid_business}}).json()["can_save"]
    cross_farm = owner.post("/assistant/execute/propose_capture", expected=422,
        json={"arguments": {**arguments, "data": {**assistant_data, "farm_id": other["id"]}}}).json()
    assert cross_farm["error"]["code"] == "farm_mismatch"
    outsider.post("/assistant/execute/check_capture", expected=404, json={"arguments": arguments})
    adviser.post("/assistant/execute/propose_capture", expected=403, headers=shared_scope, json={"arguments": arguments})
    selected_check = owner.post("/assistant/execute/check_request", json={"arguments": {
        "request_id": request["id"], "selected_operation_ids": corrected_ids}}).json()
    assert set(selected_check["check"]["selection"]["operation_ids"]) == set(corrected_ids)
    prior_count = owner.get("/operations", params={"farm_id": farm["id"]}).json()["total"]
    staged = owner.post("/assistant/execute/propose_capture", json={"arguments": arguments}).json()
    assert staged["status"] == "requires_human_review"
    assert owner.get("/operations", params={"farm_id": farm["id"]}).json()["total"] == prior_count
    with human_browser(owner, service_url) as browser:
        preview = browser.get(f"/assistant/approvals/{staged['approval_id']}")
        assert preview.status_code == 200, preview.text
        capture_check = preview.json()["payload"]["capture_check"]
        decision = {"expected_revision": staged["expected_revision"], "checksum": staged["checksum"],
            "check_token": capture_check["check_token"], "acknowledged_codes": capture_check["requires_acknowledgement"]}
        changed = browser.post(f"/assistant/approvals/{staged['approval_id']}/approve",
            json={**decision, "checksum": "0" * 64})
        assert changed.status_code == 409 and changed.json()["error"]["code"] == "proposal_changed"
        unchecked = browser.post(f"/assistant/approvals/{staged['approval_id']}/approve",
            json={**decision, "check_token": "0" * 64})
        assert unchecked.status_code == 409 and unchecked.json()["error"]["code"] == "check_changed"
        unacknowledged = browser.post(f"/assistant/approvals/{staged['approval_id']}/approve",
            json={**decision, "acknowledged_codes": []})
        assert unacknowledged.status_code == 422
        consumed = browser.post(f"/assistant/approvals/{staged['approval_id']}/approve", json=decision)
        assert consumed.status_code == 200, consumed.text
        assistant_record = owner.get(f"/operations/{consumed.json()['result']['id']}").json()
        assert assistant_record["data"]["review_status"] == "proposed"
        assert assistant_record["data"]["cz_ppp"]["efficacy"] == {"status": "pending"}
        assert owner.get("/operations", params={"farm_id": farm["id"]}).json()["total"] == prior_count + 1

        journal_arguments = {"farm_id": farm["id"], "expected_farm_revision": farm["revision"],
            "reporting_year": 2026, "operation_ids": corrected_ids, "fact_ids": [holder["id"]]}
        for change in ({"request_id": request["id"]}, {"reporting_year": 2025},
                       {"pack_id": "farm-evidence"}, {"operation_ids": []}):
            denied = owner.post("/assistant/execute/preview_journal", expected=422,
                json={"arguments": {**journal_arguments, **change}}).json()
            assert denied["error"]["code"] == "invalid_tool_arguments"
        manual_preview = owner.post("/assistant/execute/preview_journal", json={"arguments": journal_arguments}).json()
        assert manual_preview["status"] == "requires_human_review"
        manual_preview_detail = browser.get(f"/assistant/approvals/{manual_preview['approval_id']}")
        assert manual_preview_detail.status_code == 200, manual_preview_detail.text
        assert manual_preview_detail.json()["payload"]["request_id"] is None
        assert {entry["operation_id"] for entry in manual_preview_detail.json()["payload"]["snapshot"]["journals"][0]["entries"]} == set(corrected_ids)
        manual_consumed = browser.post(f"/assistant/approvals/{manual_preview['approval_id']}/approve", json={
            "expected_revision": manual_preview["expected_revision"], "checksum": manual_preview["checksum"]})
        assert manual_consumed.status_code == 200, manual_consumed.text
        manual_result = owner.get(f"/reports/{manual_consumed.json()['result']['id']}").json()
        assert manual_result["data"]["status"] == "draft" and manual_result["data"]["request_id"] is None
        assert manual_result["data"]["approval"] is None

        # A->B->C cannot be undone by rejecting B: that would revive A beside C.
        third_data = deepcopy(corrected); third_data.update(supersedes_id=replacement["id"], dose=0.7)
        third_data["cz_ppp"]["original_numbers"]["dose"] = "0.700"
        stale_stage = owner.post("/assistant/execute/propose_capture", json={"arguments": {
            **arguments, "data": third_data}}).json()
        stale_preview = browser.get(f"/assistant/approvals/{stale_stage['approval_id']}")
        assert stale_preview.status_code == 200, stale_preview.text
        stale_check = stale_preview.json()["payload"]["capture_check"]
        third, _ = capture(owner, third_data); third = review_operation(owner, third)
        rejected_ancestor = owner.post(f"/operations/{replacement['id']}/review", expected=409, json={
            "expected_revision": replacement["revision"], "decision": "rejected", "note": "Cannot revive an ancestor beside its confirmed successor."}).json()
        assert rejected_ancestor["error"]["code"] == "operation_correction_conflict"
        assert owner.get(f"/operations/{original['id']}").json()["active"] is False
        assert owner.get(f"/operations/{replacement['id']}").json()["active"] is False
        assert owner.get(f"/operations/{third['id']}").json()["active"] is True
        stale_consume = browser.post(f"/assistant/approvals/{stale_stage['approval_id']}/approve", json={
            "expected_revision": stale_stage["expected_revision"], "checksum": stale_stage["checksum"],
            "check_token": stale_check["check_token"], "acknowledged_codes": stale_check["requires_acknowledgement"]})
        assert stale_consume.status_code == 409 and stale_consume.json()["error"]["code"] == "check_changed"

    # Unselected supporting originals remain private but are still live validation
    # dependencies: stale before preparation and changed after review both block.
    for changed_before_draft in (True, False):
        marker = "UNSELECTED-PRIVATE-JOURNAL-" + uuid4().hex
        document = owner.post("/documents", expected=201, data={"farm_id": farm["id"]},
            files={"file": (marker + ".txt", (marker + "\nAuthored notebook page 14.\n").encode(), "text/plain")}).json()
        sourced_data = application(farm["id"], title="Authored source-bound application")
        sourced_data["sources"] = [{"document_id": document["id"], "document_revision": document["revision"]}]
        sourced, _ = capture(owner, sourced_data); sourced = review_operation(owner, sourced)
        selected_source = {**manual_body, "operation_ids": [sourced["id"]], "selected_document_ids": []}
        if changed_before_draft:
            owner.patch(f"/documents/{document['id']}", json={"expected_revision": document["revision"], "data": {"notes": "Reread and annotated original."}})
        source_draft = owner.post("/reports", expected=201, params=compact, json=selected_source).json()
        source_full = owner.get(f"/reports/{source_draft['id']}").json()["data"]
        assert document["id"] in source_full["validation_sources"]["journal_documents"]
        assert source_full["snapshot"]["evidence"] == []
        serialized = json.dumps(source_full["snapshot"], ensure_ascii=False)
        assert document["id"] not in serialized and marker not in serialized
        source_review = owner.get(f"/reports/{source_draft['id']}?view=review").json()["data"]
        source_binding = source_review["review"]
        approval = {"expected_revision": source_draft["revision"], "snapshot_checksum": source_binding["snapshot_checksum"],
            "check_fingerprint": source_binding["check_fingerprint"]}
        if not changed_before_draft:
            assert source_review["can_approve"]
            source_issued = owner.post(f"/reports/{source_draft['id']}/approve", json=approval).json()
            # Keep a distinct unapproved preview bound to the old source version.
            source_draft = owner.post("/reports", expected=201, params=compact, json=selected_source).json()
            source_binding = owner.get(f"/reports/{source_draft['id']}?view=review").json()["data"]["review"]
            approval = {"expected_revision": source_draft["revision"], "snapshot_checksum": source_binding["snapshot_checksum"],
                "check_fingerprint": source_binding["check_fingerprint"]}
            private_exports = {fmt: owner.get(f"/reports/{source_issued['id']}/download", params={"format": fmt}).content
                for fmt in ("json", "csv", "pdf", "xlsx")}
            assert json.loads(private_exports["json"])["snapshot"]["evidence"] == []
            assert document["id"] not in json.dumps(json.loads(private_exports["json"])["snapshot"])
            assert marker not in private_exports["csv"].decode("utf-8-sig")
            assert marker not in " ".join(page.extract_text() or "" for page in PdfReader(BytesIO(private_exports["pdf"])).pages)
            private_workbook = load_workbook(BytesIO(private_exports["xlsx"]), read_only=True, data_only=True)
            assert marker not in " ".join(str(cell) for sheet in private_workbook for row in sheet.values for cell in row)
            owner.patch(f"/documents/{document['id']}", json={"expected_revision": document["revision"], "data": {"notes": "Original reread after this preview."}})
        invalidated = owner.get(f"/reports/{source_draft['id']}?view=review").json()["data"]
        assert not invalidated["can_approve"]
        current_binding = invalidated["review"]
        blocker_page = owner.get(f"/reports/{source_draft['id']}/review", params={"section": "blockers", "offset": 0,
            "limit": 5, "expected_revision": source_draft["revision"], "snapshot_checksum": current_binding["snapshot_checksum"],
            "check_fingerprint": current_binding["check_fingerprint"]}).json()
        assert any(issue["code"] == "stale_journal_source" for issue in blocker_page["items"])
        if not changed_before_draft:
            assert invalidated["review"]["check_fingerprint"] != source_binding["check_fingerprint"]
            for fmt, content in private_exports.items():
                assert owner.get(f"/reports/{source_issued['id']}/download", params={"format": fmt}).content == content
        owner.post(f"/reports/{source_draft['id']}/approve", expected=409, json=approval)


def test_czech_journal_requestless_browser_workflow(farmers, browser_app):
    """One focused local browser case, without replaying unrelated farmer journeys."""
    owner = farmers("Browser Czech manual journal keeper")
    farm = owner.create("farms", {**farm_data("Fictional Czech journal holding"), "country": "CZ", "region": None})
    settings = {"accounts": {"czech": {"email": owner.email, "password": owner.password,
        "workspace_id": owner.workspace_id}}, "czech_farm": farm}
    directory = browser_app["directory"]
    environment = {**os.environ, "FUNKYFARM_BROWSER_URL": browser_app["url"],
        "FUNKYFARM_BROWSER_FIXTURES": json.dumps(settings),
        "FUNKYFARM_BROWSER_GREP": "Czech farmer keeps a requestless journal",
        "FUNKYFARM_BROWSER_OUTPUT": str(directory / "czech-results"),
        "FUNKYFARM_BROWSER_JUNIT": str(directory / "czech-junit.xml")}
    with (directory / "czech-playwright.log").open("w+", encoding="utf-8") as log:
        result = subprocess.run([browser_app["node"], str(PROJECT / "tests/browser/node_modules/@playwright/test/cli.js"),
            "test", "czech-journal.spec.mjs", "--config", "playwright.config.mjs"],
            cwd=PROJECT / "tests/browser", env=environment, stdout=log, stderr=subprocess.STDOUT,
            timeout=240, creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        log.seek(0)
        assert result.returncode == 0, f"Czech journal browser acceptance failed. Artifacts: {directory}\n{log.read()[-14000:]}"
