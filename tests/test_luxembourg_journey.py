"""Original LU 2026 export cells become reviewed evidence, never a filing."""

import csv
from hashlib import sha256
from io import BytesIO, StringIO
import json
import re
from uuid import uuid4

from openpyxl import load_workbook
from pypdf import PdfReader

from conftest import confirm, fact_data, farm_data


def source_csv(rows, delimiter=";", extra_last_cell=None):
    output = StringIO(newline="")
    writer = csv.writer(output, delimiter=delimiter)
    writer.writerow(list(rows[0]))
    for index, row in enumerate(rows):
        writer.writerow([row.get(key, "") for key in rows[0]] + ([extra_last_cell] if extra_last_cell and index == len(rows) - 1 else []))
    return output.getvalue().encode("utf-8")


def test_luxembourg_original_profiles_review_partial_rows_and_immutable_country_reports(farmers, tmp_path):
    owner, outsider = farmers("Luxembourg source reviewer"), farmers("Separate source reviewer")
    farm = owner.create("farms", {**farm_data("Authored Luxembourg source holding"), "country": "LU", "region": None, "farm_type": "arable"})
    foreign_farm = outsider.create("farms", {**farm_data(), "country": "LU", "region": None})
    inventory = owner.get("/imports/profiles").json()["items"]
    profiles = {item["id"]: item for item in inventory}
    assert set(profiles) == {"lu-agriculture-2026", "lu-vineyard-2026", "lu-quota-lu-2026", "lu-quota-foreign-2026"}
    for identifier, prefix, holding, identity in [
        ("lu-agriculture-2026", "lu.fa.agriculture_row.", "expl_code", ["id_parcel"]),
        ("lu-vineyard-2026", "lu.fa.vineyard_row.", "expl_code", ["id_parcel"]),
        ("lu-quota-lu-2026", "lu.fa.quota_lu_row.", "explCode", ["idParcel", "idVine"]),
        ("lu-quota-foreign-2026", "lu.fa.quota_foreign_row.", "explCode", []),
    ]:
        assert profiles[identifier]["fact_key_prefix"] == prefix
        assert profiles[identifier]["holding_column"] == holding and profiles[identifier]["identity_columns"] == identity
        assert profiles[identifier]["source_year"] == 2026 and profiles[identifier]["country"] == "LU"
    request = owner.create("requests", {"farm_id": farm["id"], "title": "Authored Luxembourg source-table preparation", "recipient": "Fictional source reviewer", "country": "LU", "region": None, "reporting_year": 2026, "deadline": None, "pack_id": "LU-cap-2026", "requirements": [], "status": "draft"})

    def source_tables(confirmed_count, proposed_count):
        guided = owner.get("/workflow/questions", params={"request_id": request["id"]}).json()
        assert guided["questions"] == [] and guided["total_count"] == 0
        tables = {table["profile_id"]: table for table in guided["source_tables"]}
        assert set(tables) == {"lu-agriculture-2026", "lu-vineyard-2026"}
        for table in tables.values():
            assert table["confirmed_count"] == confirmed_count and table["proposed_count"] == proposed_count
            assert table["label"] and table["import_url"].startswith("/app/imports?")
            assert farm["id"] in table["import_url"] and table["profile_id"] in table["import_url"]
            assert table["review_url"] == f"/app/fields/{farm['id']}#facts"
        return guided

    source_tables(0, 0)
    owner.post("/workflow/answer", expected=409, json={"request_id": request["id"], "expected_request_revision": request["revision"], "key": "lu.fa.agriculture_row.", "value": "This is not a scalar question", "unit": None, "temporal_basis": "campaign_year", "period": "2026", "note": "A row source must use its explicit import and review workflow.", "preview": False})
    assert owner.get("/facts").json()["total"] == 0
    agriculture = {"id_parcel": "000007", "field_name": "Authored north field", "expl_code": "000123", "pa_status": "N", "crop_code": "0012", "prev_crop": "0001", "a_parcel": "125,00", "a_deposit": "0", "a_eligible": "124,50", "bio": "BIO", "assur_harv": "TRUE", "activ_FLIK": "FALSE", "counter_pp": "0", "mystery_blank": "", "mystery_null": "NULL", "mystery_zero": "0", "unmapped_note": "=2+2"}
    vineyard = {"id_parcel": "000008", "Name": "Lieu-dit Étang & Coteaux", "expl_code": "000123", "pa_status": "M", "crop_code": "0013", "prev_crop": "0002", "a_parcel": "25.00", "a_quota": "20.00", "activ_FLIK": "FALSE", "Pente": "150", "Terrasse": "TRUE", "no_direct": "FALSE", "bio": "CONV_1", "corg_pourc": "2.50", "542": "ORIGINAL-CODE"}
    domestic = {"explCode": "000123", "idParcel": "000008", "idVine": "0009", "assurHarv": "FALSE", "a_quota": "12.50", "name": "Authored vineyard quota", "explMode": "Pacht", "variety": "Riesling", "aVine": "12.50", "yPlant": "2001", "grubUp": "31/12/2020", "meca": "Direktzug", "training": "Draht", "width": "180"}
    foreign = {"explCode": "000123", "country": "France", "section": "A", "idParcel": "001234", "a_quota": "0", "explMode": "Andere", "variety": "Fallow", "aVine": "0", "yPlant": "2001", "grubUp": "NULL", "meca": "Handarbeit", "training": "Pfahl", "width": "0"}

    def upload(profile, raw, *, expected=201, **overrides):
        return owner.post("/imports", expected=expected, files={"file": (profile + ".csv", raw, "text/csv")}, data={"farm_id": farm["id"], "format_profile": profile, "source_year": "2026", "delimiter": "auto", "grain": "One original authored source row", **overrides}).json()

    def review(imported, separator=".", **overrides):
        return {"expected_revision": imported["revision"], "expected_farm_revision": farm["revision"], "holding_reference": "000123", "decimal_separator": separator, "confirm_source": True, **overrides}

    def batch(imported, preview, *, acknowledge=True):
        return {"expected_revision": preview["import_revision"], "note": "Compared exact original cells, source units, identity, year and all retained warnings.", "rows": [{"row_number": item["row"]["data"]["row_number"], "expected_revision": item["row"]["revision"], "preview_token": item["preview_token"], "acknowledged_issues": item["requires_acknowledgement"] if acknowledge else []} for item in preview["items"]]}

    # Each rejected upload is atomic: neither an import nor its source rows survive.
    count = owner.get("/imports").json()["total"]
    invalid_sources = [
        (source_csv([agriculture]), {"source_year": "2025"}, "profile_source_required"),
        (source_csv([agriculture]), {"farm_id": ""}, "profile_source_required"),
        (source_csv([agriculture]), {"mapping": "[]"}, "profile_source_required"),
        (source_csv([agriculture]), {"farm_id": foreign_farm["id"]}, None),
        (source_csv([agriculture, {**agriculture, "id_parcel": "000008", "expl_code": "999999"}]), {}, "profile_mixed_holdings"),
        (source_csv([agriculture, agriculture]), {}, "profile_duplicate_identity"),
        (b"expl_code;id_parcel;id_parcel\n000123;000007;000007\n", {"delimiter": ";"}, "profile_duplicate_headers"),
        (source_csv([{**agriculture, "large_a": "x" * 200000, "large_b": "y" * 200000, "large_c": "z" * 200000}]), {}, "profile_row_too_large"),
    ]
    for raw, options, error in invalid_sources:
        result = upload("lu-agriculture-2026", raw, expected=422 if error else 404, **options)
        if error:
            assert result["error"]["code"] == error
        assert owner.get("/imports").json()["total"] == count
    assert upload("lu-quota-foreign-2026", b"explCode\n000123\n", expected=422)["error"]["code"] == "profile_delimiter_ambiguous"
    non_lu = owner.create("farms", farm_data("Not a Luxembourg source holding"))
    assert upload("lu-agriculture-2026", source_csv([agriculture]), expected=422, farm_id=non_lu["id"])["error"]["code"] == "profile_country_mismatch"

    bad = {**agriculture, "id_parcel": "000008", "a_parcel": "NaN", "assur_harv": "true", "pa_status": "Z"}
    extra = {**agriculture, "id_parcel": "000009"}
    raw = source_csv([agriculture, bad, extra], extra_last_cell="DO NOT DISCARD")
    imported = upload("lu-agriculture-2026", raw)
    outsider.get(f"/imports/{imported['id']}", expected=404)
    assert imported["data"]["sha256"] == sha256(raw).hexdigest()
    assert imported["data"]["delimiter"] == ";"
    assert owner.get(f"/imports/{imported['id']}/preview", expected=409).json()["error"]["code"] == "profile_review_required"
    assert owner.post(f"/imports/{imported['id']}/mapping", expected=422, json={"expected_revision": imported["revision"], "target": "facts", "columns": [], "defaults": {}}).json()["error"]["code"] == "profile_mapping_locked"
    for change, expected, code in [({"confirm_source": False}, 422, "profile_source_confirmation_required"), ({"holding_reference": "123"}, 422, "profile_holding_mismatch"), ({"expected_farm_revision": farm["revision"] + 1}, 409, "revision_conflict")]:
        assert owner.post(f"/imports/{imported['id']}/profile-review", expected=expected, json=review(imported, ",", **change)).json()["error"]["code"] == code
    owner.post(f"/imports/{imported['id']}/profile-review", expected=422, json=review(imported, ",", confirm_source=1))
    body = review(imported, ",")
    retry = {"Idempotency-Key": "lu-source-" + uuid4().hex}
    imported = owner.post(f"/imports/{imported['id']}/profile-review", json=body, headers=retry).json()
    assert owner.post(f"/imports/{imported['id']}/profile-review", json=body, headers=retry).json() == imported
    owner.post(f"/imports/{imported['id']}/profile-review", expected=409, json=body)
    # The source review is bound to the current farm, not merely an import checkbox.
    farm = owner.patch(f"/farms/{farm['id']}", json={"expected_revision": farm["revision"], "data": {"notes": "Source review must use this changed farm revision."}}).json()
    owner.get(f"/imports/{imported['id']}/preview", expected=409)
    imported = owner.post(f"/imports/{imported['id']}/profile-review", json=review(imported, ",")).json()
    preview = owner.get(f"/imports/{imported['id']}/preview").json()
    assert preview["items"][0]["data"]["value"] == agriculture
    assert {"profile_invalid_number", "profile_invalid_boolean", "profile_invalid_code"} <= {issue["code"] for issue in preview["items"][1]["issues"]}
    assert "profile_source_column_mismatch" in {issue["code"] for issue in preview["items"][2]["issues"]}
    assert {"profile_unknown_columns", "profile_previous_year_crop", "profile_bio_source_only", "profile_area_arithmetic_review"} <= set(preview["items"][0]["requires_acknowledgement"])
    first = preview["items"][0]
    owner.post(f"/imports/{imported['id']}/rows/1/preview", expected=422, json={"expected_revision": first["row"]["revision"], "target": "facts", "data": {**first["data"], "value": {**agriculture, "a_parcel": "0"}}})
    unacknowledged = owner.post(f"/imports/{imported['id']}/rows/review", json=batch(imported, preview, acknowledge=False)).json()
    assert unacknowledged["saved_count"] == 0
    assert unacknowledged["items"][0]["error"]["code"] == "review_warnings_required"
    result = owner.post(f"/imports/{imported['id']}/rows/review", json=batch(imported, preview)).json()
    assert (result["saved_count"], result["failed_count"]) == (1, 2)
    saved = [result["items"][0]["record"]]
    repeat = owner.client.post(f"/imports/{imported['id']}/rows/review", json=batch(imported, preview))
    assert repeat.status_code in (200, 409)
    assert owner.get("/facts").json()["total"] == 1
    originals = owner.get(f"/imports/{imported['id']}/rows").json()["items"]
    assert originals[2]["data"]["original_row"][-1] == "DO NOT DISCARD"
    assert originals[1]["data"]["review_status"] == "unreviewed"
    sources = [(imported, agriculture, ",")]

    # Same parcel, different quota subparcel is a distinct documented identity.
    scenarios = [("lu-vineyard-2026", [vineyard, {**vineyard, "id_parcel": "000010", "a_parcel": "-1", "corg_pourc": "150"}], "\t"),
                 ("lu-quota-lu-2026", [domestic, {**domestic, "idVine": "0010"}], ","),
                 ("lu-quota-foreign-2026", [foreign, foreign], ";")]
    for profile_id, rows, delimiter in scenarios:
        record = upload(profile_id, source_csv(rows, delimiter))
        assert record["data"]["delimiter"] == delimiter
        record = owner.post(f"/imports/{record['id']}/profile-review", json=review(record)).json()
        check = owner.get(f"/imports/{record['id']}/preview").json()
        if profile_id == "lu-vineyard-2026":
            assert check["items"][0]["can_save"] is True, "A 150% slope is not a 150% composition."
            assert {"profile_negative_quantity", "profile_percentage_out_of_range"} <= {issue["code"] for issue in check["items"][1]["issues"]}
        for index, row in enumerate(rows):
            if profile_id == "lu-vineyard-2026" and index == 1:
                continue
            check = owner.get(f"/imports/{record['id']}/preview").json()
            check["items"] = [check["items"][index]]
            outcome = owner.post(f"/imports/{record['id']}/rows/review", json=batch(record, check)).json()
            assert (outcome["saved_count"], outcome["failed_count"]) == (1, 0)
            saved.append(outcome["items"][0]["record"])
            sources.append((record, row, "."))
    assert len({item["data"]["key"] for item in saved}) == 6
    source_tables(0, 1)
    for fact, (source, row, decimal) in zip(saved, sources):
        profile = profiles[source["data"]["format_profile"]]
        assert re.fullmatch(re.escape(profile["fact_key_prefix"]) + r"[a-f0-9]{64}", fact["data"]["key"])
        assert fact["data"]["value"] == row and all(isinstance(value, str) for value in fact["data"]["value"].values())
        assert fact["data"]["review_status"] == "proposed" and fact["data"]["unit"] is None
        assert fact["data"]["period"] == "2026" and fact["data"]["temporal_basis"] == "campaign_year"
        assert not fact["data"].get("parcel_id")
        provenance = next(entry for entry in fact["data"]["sources"] if entry.get("import_id"))
        assert provenance["import_id"] == source["id"] and provenance["sha256"] == source["data"]["sha256"]
        assert provenance["source_name"] == source["data"]["source_name"] and provenance["grain"] == "One original authored source row"
        assert provenance["row_id"] and provenance["row_number"] in (1, 2)
        assert any(entry.get("adapter_profile") == profile["id"] and entry.get("decimal_separator") == decimal for entry in fact["data"]["sources"])
    assert owner.get(f"/farms/{farm['id']}").json() == farm
    assert owner.get("/operations").json()["total"] == owner.get("/purchases").json()["total"] == 0

    # An empty mapped template and unreviewed rows cannot be issued.
    for ids in ([], [saved[0]["id"]]):
        draft = owner.post("/reports", expected=201, json={"farm_id": farm["id"], "pack_id": "LU-cap-2026", "reporting_year": 2026, "fact_ids": ids}).json()
        denied = owner.post(f"/reports/{draft['id']}/approve", expected=409, json={"expected_revision": draft["revision"]}).json()
        if not ids:
            assert "mapped_rows_required" in json.dumps(denied)
    saved = [confirm(owner, item, "Checked this original 2026 row and its raw source warnings; not an authority acceptance.") for item in saved]
    source_tables(1, 0)
    retained = []
    for pack_id, selected in [("LU-cap-2026", saved[:2]), ("LU-annual-2026", saved[2:])]:
        draft = owner.post("/reports", expected=201, json={"farm_id": farm["id"], "pack_id": pack_id, "reporting_year": 2026, "fact_ids": [item["id"] for item in selected]}).json()
        issued = owner.post(f"/reports/{draft['id']}/approve", json={"expected_revision": draft["revision"]}).json()
        assert issued["data"]["status"] == "issued" and issued["data"]["filing_status"] == "not_filed"
        fields = issued["data"]["snapshot"]["fields"]
        for fact in selected:
            field = next(item for item in fields if item.get("fact_id") == fact["id"])
            assert field["fact_key"] == fact["data"]["key"] and field["value"] == fact["data"]["value"]
            schema = field["row_schema"]
            assert schema["id"] in profiles and schema["holding_column"] in field["value"]
            assert schema["identity_columns"] == profiles[schema["id"]]["identity_columns"]
            assert schema["row_columns"] == profiles[schema["id"]]["row_columns"]
            provenance = next(source for source in field["sources"] if source.get("import_id"))
            original = next(source for source in fact["data"]["sources"] if source.get("import_id"))
            for key in ("import_id", "row_id", "row_number", "sha256", "source_name", "grain"):
                assert provenance[key] == original[key]
            assert any(source.get("adapter_profile") == schema["id"] and source.get("source_year") == 2026 for source in field["sources"])
        exports = {}
        for format in ("json", "csv", "xlsx", "pdf"):
            response = owner.get(f"/reports/{issued['id']}/download?format={format}")
            outsider.get(f"/reports/{issued['id']}/download?format={format}", expected=404)
            assert "no-store" in response.headers["cache-control"]
            exports[format] = response.content
        assert json.loads(exports["json"])["snapshot"] == issued["data"]["snapshot"]
        rows = list(csv.DictReader(StringIO(exports["csv"].decode("utf-8-sig"))))
        workbook = load_workbook(BytesIO(exports["xlsx"]), data_only=False)
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        xlsx_rows = [dict(zip(headers, row)) for row in sheet.iter_rows(min_row=2, values_only=True)]
        assert all(cell.data_type != "f" for row in sheet for cell in row)
        for fact in selected:
            aggregate = next(row for row in rows if row["fact_id"] == fact["id"] and json.loads(row["value_json"]) == fact["data"]["value"])
            assert json.loads(aggregate["value_json"]) == fact["data"]["value"]
            for code, value in fact["data"]["value"].items():
                cell = next(row for row in rows if row["field"] == fact["data"]["key"] + "." + code)
                assert json.loads(cell["value_json"]) == value
                assert any(row["field"] == cell["field"] and row["value_json"] == cell["value_json"] for row in xlsx_rows)
                if code in ("a_parcel", "a_deposit", "a_eligible", "a_quota", "aVine", "width", "Pente", "corg_pourc"):
                    expected_unit = "cm" if code == "width" else "%" if code in ("Pente", "corg_pourc") else "a"
                    assert json.loads(cell["unit_json"]) == expected_unit
        reader = PdfReader(BytesIO(exports["pdf"]))
        # ReportLab's invariant epoch must replace the render wall clock, so
        # an unchanged issued snapshot remains byte-identical on another day.
        assert reader.metadata.get("/CreationDate", "").startswith("D:20000101")
        pdf = "\n".join(page.extract_text() for page in reader.pages)
        assert "000123" in pdf and "not" in pdf.lower()
        if pack_id == "LU-cap-2026":
            area = next(row for row in rows if row["field"] == saved[0]["data"]["key"] + ".a_parcel")
            assert "Measured parcel area" in area["label"] and "a_parcel" in area["label"] and json.loads(area["unit_json"]) == "a"
            unknown = next(row for row in rows if row["field"].endswith(".mystery_blank"))
            assert json.loads(unknown["unit_json"]) is None and json.loads(unknown["value_json"]) == ""
            assert "Measured parcel area" in pdf and "a_parcel" in pdf and "125,00" in pdf
        else:
            assert "Quota area" in pdf and "a_quota" in pdf and "180" in pdf
        (tmp_path / (pack_id + ".pdf")).write_bytes(exports["pdf"])
        retained.append((issued, exports))
    correction = owner.create("facts", {key: value for key, value in saved[0]["data"].items() if key not in ("reviewed_at", "reviewed_by", "review_note")} | {"supersedes_id": saved[0]["id"], "review_status": "proposed", "value": {**agriculture, "a_parcel": "126,00"}})
    confirm(owner, correction, "Explicit corrected source observation; the issued source snapshot remains retained.")
    for issued, exports in retained:
        for format, content in exports.items():
            assert owner.get(f"/reports/{issued['id']}/download?format={format}").content == content

    # Legitimate bounded source rows can produce a report above both native
    # transport limits. Use the real upload/review path, not injected snapshots.
    farm = owner.create("farms", {**farm_data("Large source-table review holding"), "country": "LU", "region": None})
    large_rows = [{**agriculture, "id_parcel": f"{100000 + index:06d}", "field_name": f"Source field {index + 1}",
                   **{f"retained_note_{part}": f"Original row {index + 1}, note {part}. " + chr(65 + part) * 149950 for part in range(3)}}
                  for index in range(32)]
    imported = upload("lu-agriculture-2026", source_csv(large_rows))
    imported = owner.post(f"/imports/{imported['id']}/profile-review", json=review(imported, ",")).json()
    large_facts = []
    for offset in range(0, len(large_rows), 5):
        preview = owner.get(f"/imports/{imported['id']}/preview", params={"offset": offset, "limit": 5}).json()
        outcome = owner.post(f"/imports/{imported['id']}/rows/review", json=batch(imported, preview)).json()
        assert outcome["failed_count"] == 0 and outcome["saved_count"] == len(preview["items"])
        large_facts.extend(confirm(owner, item["record"], "Checked every original source cell, unit, period and warning.") for item in outcome["items"])
    assert len(large_facts) == len(large_rows)
    evidence = owner.post("/documents", expected=201, files={"file": ("selected-source.txt", b"Fictional reviewed source evidence.\n", "text/plain")}, data={"farm_id": farm["id"]}).json()
    activity = owner.create("operations", {"farm_id": farm["id"], "title": "Recorded field visit", "type": "observation", "date": "2026-06-01", "period": "2026-06-01", "temporal_basis": "calendar_date", "review_status": "proposed"})
    sale = owner.create("sales", {"farm_id": farm["id"], "title": "Fictional vegetable sale", "category": "produce", "date": "2026-06-01", "buyer": "Fictional local buyer", "quantity": 25, "unit": "kg", "amount": 50, "currency": "EUR", "review_status": "proposed"})
    for kind, record in (("operations", activity), ("sales", sale)):
        reviewed = owner.post(f"/{kind}/{record['id']}/review", json={"expected_revision": record["revision"], "decision": "confirmed", "note": "Checked this actual dated fictional ledger record; not a derived farm activity."}).json()
        record.update(reviewed)
    linked = confirm(owner, owner.create("facts", fact_data(farm["id"], key="reviewed_sale_reference", value="Fictional sale reference", unit=None,
                     sources=[{"record_id": item["id"], "record_kind": kind, "revision": item["revision"]} for kind, item in (("sales", sale), ("operations", activity))])))
    body = {"farm_id": farm["id"], "pack_id": "LU-cap-2026", "reporting_year": 2026,
            "fact_ids": [item["id"] for item in [*large_facts, linked]], "selected_document_ids": [evidence["id"]], "operation_ids": [activity["id"]]}
    compact = {"view": "summary"}
    retry = {"Idempotency-Key": "large-report-" + uuid4().hex}
    response = owner.post("/reports", expected=201, params=compact, headers=retry, json=body)
    draft = response.json()
    assert len(response.content) < 64 * 1024 and draft["data"]["transport_view"] == "summary"
    assert not {"fields", "facts", "operations", "evidence"} & draft["data"]["snapshot"].keys()
    assert "validation_sources" not in draft["data"]
    assert owner.post("/reports", expected=201, params=compact, headers=retry, json=body).json() == draft
    legacy = owner.post("/reports", expected=201, headers=retry, json=body)
    assert legacy.json()["id"] == draft["id"] and "transport_view" not in legacy.json()["data"], "Changing response view must replay the same committed creation."
    assert len(legacy.content) > 24 * 1024 * 1024, "Fixture must exceed both the old12MiB JSON and24MiB download caps."
    retained_snapshot = legacy.json()["data"]["snapshot"]
    metadata = owner.get(f"/reports/{draft['id']}", params={"view": "review"}).json()
    assert len(json.dumps(metadata).encode()) < 64 * 1024
    assert metadata["data"]["transport_view"] == "review" and metadata["data"]["can_approve"] is True
    binding = metadata["data"]["review"]
    assert binding["sections"] == {"fields": 33, "operations": 1, "evidence": 1, "blockers": 0, "changes": 0, "receipts": 0, "reissues": 0}
    assert binding["page_limit"] == 5 and binding["byte_limit"] == 2 * 1024 * 1024

    def page_params(binding, section, offset=0, **overrides):
        return {"section": section, "expected_revision": binding["report_revision"], "snapshot_checksum": binding["snapshot_checksum"],
                "check_fingerprint": binding["check_fingerprint"], "offset": offset, "limit": 5, **overrides}

    def read_section(report_id, binding, section):
        items, offset = [], 0
        reduced_page = False
        while True:
            response = owner.get(f"/reports/{report_id}/review", params=page_params(binding, section, offset))
            page = response.json()
            assert len(response.content) <= binding["byte_limit"]
            assert page["section"] == section and page["offset"] == offset and page["total"] == binding["sections"][section]
            assert page["report_revision"] == binding["report_revision"] and page["snapshot_checksum"] == binding["snapshot_checksum"]
            assert page["check_fingerprint"] == binding["check_fingerprint"]
            assert len(page["items"]) <= 5
            items.extend(page["items"])
            if page["next_offset"] is None:
                assert len(items) == page["total"]
                break
            assert 0 < len(page["items"]) and page["next_offset"] == offset + len(page["items"])
            reduced_page |= len(page["items"]) < 5
            offset = page["next_offset"]
        return items, reduced_page

    fields, reduced = read_section(draft["id"], binding, "fields")
    assert reduced, "The byte bound must stop before five complete450KiB source rows."
    assert fields == retained_snapshot["fields"]
    assert [field["value"] for field in fields[:32]] == large_rows
    assert len({field["fact_id"] for field in fields}) == 33
    for field in fields[:32]:
        assert field["row_schema"]["id"] == "lu-agriculture-2026" and field["row_schema"]["row_columns"] == profiles["lu-agriculture-2026"]["row_columns"]
        assert any(source.get("import_id") == imported["id"] and source.get("sha256") == imported["data"]["sha256"] for source in field["sources"])
    for section in ("operations", "evidence", "blockers", "changes", "receipts", "reissues"):
        items, _ = read_section(draft["id"], binding, section)
        if section in ("operations", "evidence"):
            assert items == retained_snapshot[section]
    outsider.get(f"/reports/{draft['id']}?view=review", expected=404)
    outsider.get(f"/reports/{draft['id']}/review", params=page_params(binding, "fields"), expected=404)
    owner.get(f"/reports/{draft['id']}/review", params=page_params(binding, "fields", expected_revision=binding["report_revision"] + 1), expected=409)
    owner.get(f"/reports/{draft['id']}/review", params=page_params(binding, "fields", snapshot_checksum="0" * 64), expected=409)
    owner.get(f"/reports/{draft['id']}/review", params=page_params(binding, "blockers", check_fingerprint="0" * 64), expected=409)
    owner.get(f"/reports/{draft['id']}/review", params=page_params(binding, "fields", limit=6), expected=422)
    approval = {"expected_revision": draft["revision"], "snapshot_checksum": binding["snapshot_checksum"], "check_fingerprint": binding["check_fingerprint"]}
    owner.post(f"/reports/{draft['id']}/approve", params=compact, json={**approval, "check_fingerprint": "0" * 64}, expected=409)
    approve_retry = {"Idempotency-Key": "large-approve-" + uuid4().hex}
    issued_response = owner.post(f"/reports/{draft['id']}/approve", params=compact, json=approval, headers=approve_retry)
    issued = issued_response.json()
    assert len(issued_response.content) < 64 * 1024 and issued["data"]["status"] == "issued"
    assert owner.post(f"/reports/{draft['id']}/approve", params=compact, json=approval, headers=approve_retry).json() == issued
    for kind, record in (("facts", large_facts[0]), ("operations", activity), ("sales", sale)):
        dependencies = owner.get(f"/{kind}/{record['id']}/reports", params=compact)
        assert len(dependencies.content) < 64 * 1024 and {item["id"] for item in dependencies.json()["items"]} == {issued["id"]}
        assert all(item["data"]["transport_view"] == "summary" for item in dependencies.json()["items"])
        outsider.get(f"/{kind}/{record['id']}/reports", params=compact, expected=404)
    summary_list = owner.get("/reports", params={**compact, "farm_id": farm["id"]})
    assert len(summary_list.content) < 64 * 1024 and summary_list.json()["total"] == 1
    original_hashes = {}
    for format in ("json", "csv", "xlsx"):
        response = owner.get(f"/reports/{issued['id']}/download", params={"format": format})
        original_hashes[format] = sha256(response.content).hexdigest()
        if format == "json":
            assert len(response.content) > 24 * 1024 * 1024 and response.json()["snapshot"] == retained_snapshot
        assert response.headers["X-Report-Snapshot-SHA256"] == binding["snapshot_checksum"]
    pending = owner.post("/reports", expected=201, params=compact, json=body).json()
    pending_binding = owner.get(f"/reports/{pending['id']}", params={"view": "review"}).json()["data"]["review"]
    old_review = owner.get(f"/reports/{issued['id']}", params={"view": "review"}).json()["data"]["review"]
    correction = owner.create("facts", {key: value for key, value in large_facts[0]["data"].items() if key not in ("reviewed_at", "reviewed_by", "review_note")} |
                              {"supersedes_id": large_facts[0]["id"], "review_status": "proposed", "value": {**large_rows[0], "a_parcel": "126,00"}})
    confirm(owner, correction, "Checked separate corrected observation; retained report bytes must not change.")
    changed = owner.get(f"/reports/{issued['id']}", params={"view": "review"}).json()
    changed_binding = changed["data"]["review"]
    assert changed["data"]["amendment"]["needed"] and changed_binding["sections"]["changes"] > 0
    stale_approval = owner.post(f"/reports/{pending['id']}/approve", expected=409, params=compact,
                               json={"expected_revision": pending["revision"], "snapshot_checksum": pending_binding["snapshot_checksum"], "check_fingerprint": pending_binding["check_fingerprint"]})
    assert stale_approval.json()["error"]["code"] == "review_checks_changed" and len(stale_approval.content) < 64 * 1024
    blocked_approval = owner.post(f"/reports/{pending['id']}/approve", expected=409, params=compact, json={"expected_revision": pending["revision"]})
    assert blocked_approval.json()["error"]["code"] == "approval_blocked"
    assert blocked_approval.json()["error"]["details"]["blocker_count"] > 0 and "blockers" not in blocked_approval.json()["error"]["details"]
    owner.get(f"/reports/{issued['id']}/review", params=page_params(old_review, "changes"), expected=409)
    assert read_section(issued["id"], changed_binding, "changes")[0]
    assert read_section(issued["id"], changed_binding, "fields")[0] == retained_snapshot["fields"]
    reissue_body = {"expected_revision": issued["revision"], "snapshot_checksum": issued["data"]["checksum"]}
    reissue_retry = {"Idempotency-Key": "large-reissue-" + uuid4().hex}
    reissued = owner.post(f"/reports/{issued['id']}/reissue", expected=201, params=compact, headers=reissue_retry, json=reissue_body).json()
    assert reissued["data"]["correction_of"] == issued["id"] and reissued["data"]["transport_view"] == "summary"
    assert owner.post(f"/reports/{issued['id']}/reissue", expected=201, params=compact, headers=reissue_retry, json=reissue_body).json() == reissued
    owner.post(f"/reports/{issued['id']}/receipts", expected=201, json={**reissue_body, "reference": "FICTIONAL-MANUAL-LARGE-REPORT", "date": "2026-10-17"})
    final_binding = owner.get(f"/reports/{issued['id']}", params={"view": "review"}).json()["data"]["review"]
    assert read_section(issued["id"], final_binding, "receipts")[0][0]["data"]["authority_acceptance_verified"] is False
    assert read_section(issued["id"], final_binding, "reissues")[0][0]["id"] == reissued["id"]
    for format, original_hash in original_hashes.items():
        assert sha256(owner.get(f"/reports/{issued['id']}/download", params={"format": format}).content).hexdigest() == original_hash
